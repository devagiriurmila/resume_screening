import io

import openpyxl

from services.report_service import build_csv, safe_cell
from tests.conftest import JOB


def test_excel_report_download(client, uploaded):
    sid = uploaded["session_id"]
    r = client.get(f"/api/reports/excel?session_id={sid}")
    assert r.status_code == 200
    assert "spreadsheetml" in r.mimetype
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "Candidates", "Skill Breakdown"]
    ws = wb["Candidates"]
    header = [c.value for c in ws[1]]
    assert "Skill Match %" in header and ws.max_row == 2
    row = dict(zip(header, [c.value for c in ws[2]]))
    assert row["Candidate"] == "Test Person" and row["Skill Match %"] == 100
    assert "Job" in header and row["Job"] == JOB["title"] and row["Screened At"]
    assert wb["Skill Breakdown"].max_row == 5  # header + 4 required skills
    notice = " ".join(str(c.value) for c in wb["Summary"]["B"])
    assert "human recruiter" in notice


def test_excel_report_all_sessions_and_empty(client, job_id):
    assert client.get("/api/reports/excel").status_code == 404  # nothing yet
    assert client.get("/api/reports/excel?session_id=999").status_code == 404


def test_csv_report(client, uploaded):
    r = client.get(f"/api/reports/csv?session_id={uploaded['session_id']}")
    assert r.status_code == 200 and r.mimetype == "text/csv"
    text = r.get_data(as_text=True)
    assert "Test Person" in text and "Skill Match %" in text


def test_summary_report(client, uploaded):
    r = client.get(f"/api/reports/summary?session_id={uploaded['session_id']}")
    body = r.get_data(as_text=True)
    assert r.status_code == 200 and "SCREENING SUMMARY" in body and "Test Analyst" in body
    assert "must not be used to reject" in body


def test_formula_injection_is_neutralised():
    assert safe_cell("=HYPERLINK(\"http://x\")").startswith("'")
    assert safe_cell("@SUM(A1)").startswith("'")
    assert safe_cell("-cmd|' /C calc'!A0").startswith("'")
    assert safe_cell("+91 90000 00101") == "+91 90000 00101"  # real phone numbers stay readable
    assert safe_cell("Normal text") == "Normal text"


def test_csv_escapes_formula_names():
    rows = [{"candidate_id": 1, "name": "=1+1", "email": None, "phone": None, "education": None, "experience_years": None,
             "experience_method": None, "skills": [], "score": 10.0, "matched_skills": [], "missing_skills": ["Python"],
             "needs_review": True, "review_reasons": [], "review_status": "pending", "reviewer_note": "",
             "processing_status": "processed", "original_filename": "a.pdf", "session_id": 1, "job_title": "J",
             "screened_at": "2026-01-01 00:00:00"}]
    out = build_csv(rows)
    assert "'=1+1" in out and "Not detected" in out
