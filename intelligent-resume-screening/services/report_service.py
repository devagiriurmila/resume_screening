"""Excel / CSV / text summary reports. Excel is written with pandas + openpyxl."""
import csv
import io
import re
from datetime import datetime
from typing import Dict, List, Optional

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DISCLAIMER = ("Scores show only how many recruiter-defined required skills appear in each resume. They are not a "
              "measure of overall candidate quality and must not be used to reject anyone automatically. "
              "A human recruiter must review every result.")
NOT_DETECTED = "Not detected"
REVIEW_LABELS = {"pending": "Pending review", "shortlisted": "Shortlisted", "on_hold": "On hold",
                 "not_selected": "Not selected (human decision)"}


PHONE_LIKE = re.compile(r"^\+?[\d\s().\-]{7,}$")


def safe_cell(value):
    """Neutralise spreadsheet formula injection (cells starting with = + - @), but keep plain phone numbers."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@") and not PHONE_LIKE.match(value):
        return "'" + value
    return value


def _experience(row: Dict) -> str:
    if row.get("experience_years") is None:
        return NOT_DETECTED
    return f"{row['experience_years']} ({row.get('experience_method') or 'detected'})"


def candidate_table(rows: List[Dict]) -> pd.DataFrame:
    ordered = sorted(rows, key=lambda r: (-r["score"], (r.get("name") or "~").lower()))
    data = []
    for rank, r in enumerate(ordered, start=1):
        data.append({
            "Rank": rank,
            "Candidate": r.get("name") or NOT_DETECTED,
            "Email": r.get("email") or NOT_DETECTED,
            "Phone": r.get("phone") or NOT_DETECTED,
            "Experience (years)": _experience(r),
            "Education": r.get("education") or NOT_DETECTED,
            "Skill Match %": r["score"],
            "Matched Skills": ", ".join(r["matched_skills"]) or "None",
            "Missing / Unconfirmed Skills": ", ".join(r["missing_skills"]) or "None",
            "All Detected Skills": ", ".join(r["skills"]) or "None",
            "Processing Status": r["processing_status"].replace("_", " "),
            "Needs Manual Review": "Yes" if r["needs_review"] else "No",
            "Review Reasons": " | ".join(r["review_reasons"]),
            "Human Review Status": REVIEW_LABELS.get(r["review_status"], r["review_status"]),
            "Reviewer Note": r.get("reviewer_note") or "",
            "Source File": r.get("original_filename") or "",
            "Screening Session": r["session_id"],
            "Job": r["job_title"],
            "Screened At": r["screened_at"],
        })
    return pd.DataFrame(data)


def breakdown_table(rows: List[Dict], breakdowns: Dict[int, List[dict]]) -> pd.DataFrame:
    data = []
    for r in rows:
        for b in breakdowns.get(r["candidate_id"], []):
            data.append({"Candidate": r.get("name") or NOT_DETECTED, "Required Skill": b["skill"],
                         "Weight %": b["weight"], "Found": "Yes" if b["matched"] else "No",
                         "Points": b["points"], "Evidence": b.get("evidence") or ""})
    return pd.DataFrame(data, columns=["Candidate", "Required Skill", "Weight %", "Found", "Points", "Evidence"])


def summary_stats(rows: List[Dict], high_threshold: float) -> Dict:
    n = len(rows)
    return {
        "count": n,
        "average": round(sum(r["score"] for r in rows) / n, 1) if n else 0,
        "high": sum(1 for r in rows if r["score"] >= high_threshold),
        "review": sum(1 for r in rows if r["needs_review"]),
    }


def build_excel(rows: List[Dict], breakdowns: Dict[int, List[dict]], path: str, title: str,
                high_threshold: float = 70) -> str:
    stats = summary_stats(rows, high_threshold)
    cand = candidate_table(rows).map(safe_cell)
    brk = breakdown_table(rows, breakdowns).map(safe_cell)
    summary = pd.DataFrame([
        ("Report", title), ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Resumes in report", stats["count"]), ("Average skill-match %", stats["average"]),
        (f"Candidates with score >= {high_threshold:g}%", stats["high"]),
        ("Flagged for manual review", stats["review"]), ("Important notice", DISCLAIMER),
    ], columns=["Item", "Value"])

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        cand.to_excel(xw, sheet_name="Candidates", index=False)
        brk.to_excel(xw, sheet_name="Skill Breakdown", index=False)
        head_fill = PatternFill("solid", fgColor="0D6EFD")
        for name, ws in xw.sheets.items():
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = head_fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            if name != "Summary" and ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for idx, col in enumerate(ws.columns, start=1):
                longest = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 60)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.alignment = Alignment(vertical="top", wrap_text=True)
        xw.sheets["Summary"].column_dimensions["B"].width = 90
    return path


def build_csv(rows: List[Dict]) -> str:
    df = candidate_table(rows).map(safe_cell)
    buf = io.StringIO()
    df.to_csv(buf, index=False, quoting=csv.QUOTE_NONNUMERIC)
    return buf.getvalue()


def build_summary_text(rows: List[Dict], title: str, high_threshold: float, session: Optional[Dict] = None) -> str:
    stats = summary_stats(rows, high_threshold)
    ordered = sorted(rows, key=lambda r: -r["score"])
    lines = ["SCREENING SUMMARY", "=" * 60, f"Report      : {title}",
             f"Generated   : {datetime.now():%Y-%m-%d %H:%M:%S}"]
    if session:
        lines += [f"Session     : #{session['id']} ({session['created_at']})", f"Job         : {session['job_title']}"]
        weights = session.get("job_snapshot", {}).get("skills", [])
        if weights:
            lines.append("Skill weights: " + ", ".join(f"{w['skill']} {w['weight']:g}%" for w in weights))
    lines += ["", f"Resumes processed        : {stats['count']}", f"Average skill-match      : {stats['average']}%",
              f"Score >= {high_threshold:g}%            : {stats['high']}",
              f"Flagged for manual review: {stats['review']}", "", "TOP RESULTS (for human review)", "-" * 60]
    for i, r in enumerate(ordered[:10], start=1):
        lines.append(f"{i:>2}. {r.get('name') or NOT_DETECTED:<28} {r['score']:>6.2f}%  matched: "
                     f"{', '.join(r['matched_skills']) or 'none'}")
    lines += ["", "NOTICE", "-" * 60, DISCLAIMER, ""]
    return "\n".join(lines)
